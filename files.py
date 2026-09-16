"""
Generación de archivos: Word, PowerPoint y páginas web (HTML) simples,
a partir de un tema que pide el alumno. Usa la IA para armar el contenido
y estas funciones para darle el formato final al archivo.
"""
import os
import re
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "generados")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def _safe_filename(texto):
    texto = re.sub(r"[^a-zA-Z0-9_\- ]", "", texto).strip().replace(" ", "_")
    return texto[:40] or "documento"


def crear_word(tema: str, contenido: str) -> str:
    """contenido: texto con títulos y párrafos, separados por líneas.
    Las líneas que empiezan con '# ' se tratan como títulos."""
    doc = Document()
    doc.add_heading(tema, level=0)

    for linea in contenido.split("\n"):
        linea = linea.strip()
        if not linea:
            continue
        if linea.startswith("# "):
            doc.add_heading(linea[2:].strip(), level=1)
        elif linea.startswith("- ") or linea.startswith("* "):
            doc.add_paragraph(linea[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(linea)

    path = os.path.join(OUTPUT_DIR, f"{_safe_filename(tema)}.docx")
    doc.save(path)
    return path


def crear_powerpoint(tema: str, diapositivas: list) -> str:
    """diapositivas: lista de dicts {"titulo": str, "puntos": [str, ...]}"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide de portada
    layout_portada = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout_portada)
    slide.shapes.title.text = tema
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Generado por Ayudante Escolar"

    layout_contenido = prs.slide_layouts[1]
    for d in diapositivas:
        slide = prs.slides.add_slide(layout_contenido)
        slide.shapes.title.text = d.get("titulo", "")
        body = slide.placeholders[1].text_frame
        body.clear()
        puntos = d.get("puntos", [])
        for i, punto in enumerate(puntos):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = punto
            p.font.size = PptPt(20)

    path = os.path.join(OUTPUT_DIR, f"{_safe_filename(tema)}.pptx")
    prs.save(path)
    return path


def crear_pagina_web(tema: str, html_body: str) -> str:
    """html_body: el contenido HTML del cuerpo de la página (lo genera la IA)."""
    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<title>{tema}</title>
<style>
  body {{ font-family: Arial, sans-serif; max-width: 800px; margin: 40px auto;
          padding: 0 20px; line-height: 1.6; color: #222; }}
  h1, h2 {{ color: #1a3c6e; }}
  ul {{ padding-left: 20px; }}
</style>
</head>
<body>
{html_body}
</body>
</html>"""
    path = os.path.join(OUTPUT_DIR, f"{_safe_filename(tema)}.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    return path
